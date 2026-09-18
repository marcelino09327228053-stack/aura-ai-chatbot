"""Core user-journey tests using an isolated temporary SQLite database."""

import os
import sys
import tempfile
import unittest
import asyncio
import hashlib
import hmac
import json
import re
from urllib.parse import parse_qs, urlparse
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

test_db = Path(tempfile.gettempdir()) / "aura_core_flow_test.db"
test_db.unlink(missing_ok=True)

os.environ["SQLITE_PATH"] = str(test_db)
os.environ["SECRET_KEY"] = "test-secret-key-that-is-at-least-32-bytes"
os.environ["REDIS_ENABLED"] = "false"
os.environ["AGENTS_ENABLED"] = "false"
os.environ["PLUGINS_ENABLED"] = "false"
os.environ["GEMINI_API_KEY"] = ""
os.environ["AURA_ENV"] = "development"
os.environ["SMTP_HOST"] = ""
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient

from main import app


class CoreFlowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from app.infrastructure.database import sqlite_backend
        from app.database.connection import init_db
        if sqlite_backend._conn is not None:
            sqlite_backend._conn.close()
            sqlite_backend._conn = None
        sqlite_backend.SQLITE_PATH = str(test_db)
        test_db.unlink(missing_ok=True)
        init_db()
        cls.client = TestClient(app, raise_server_exceptions=False)
        response = cls.client.post(
            "/auth/register",
            json={
                "email": "core-flow@example.com",
                "password": "test-password",
                "company_name": "Core Flow Company",
            },
        )
        assert response.status_code == 200, response.text
        cls.token = response.json()["token"]
        cls.headers = {"Authorization": f"Bearer {cls.token}"}
        from app.services.subscription_service import activate_paid_cycle
        activate_paid_cycle(response.json()["company"]["id"], "test")

    @classmethod
    def tearDownClass(cls):
        from app.infrastructure.database import sqlite_backend

        if sqlite_backend._conn is not None:
            sqlite_backend._conn.close()
            sqlite_backend._conn = None
        test_db.unlink(missing_ok=True)

    def test_authentication_and_company_profile(self):
        health = self.client.get("/health")
        self.assertEqual(health.status_code, 200)
        self.assertEqual(health.json()["database"], "ok")
        self.assertEqual(health.headers["x-content-type-options"], "nosniff")
        for page in ("/dashboard", "/widget", "/knowledge", "/support", "/team", "/crm"):
            self.assertEqual(self.client.get(page).status_code, 200)

        response = self.client.get("/auth/me", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        company_id = response.json()["company"]["id"]

        update = self.client.put(
            f"/companies/{company_id}",
            headers=self.headers,
            json={"company_profile": "A test company that sells software."},
        )
        self.assertEqual(update.status_code, 200)
        self.assertIn("sells software", update.json()["company_profile"])

        account = self.client.put(
            "/auth/account",
            headers=self.headers,
            json={"full_name": "Aura Test Customer"},
        )
        self.assertEqual(account.status_code, 200)
        self.assertEqual(account.json()["user"]["full_name"], "Aura Test Customer")

    def test_dashboard_creation_switch_and_data_isolation(self):
        registration = self.client.post(
            "/auth/register",
            json={
                "email": "workspace-owner@example.com",
                "password": "test-password",
                "company_name": "Workspace One",
            },
        )
        self.assertEqual(registration.status_code, 200)
        first_token = registration.json()["token"]
        first_headers = {"Authorization": f"Bearer {first_token}"}
        first_company = registration.json()["company"]

        from app.services.subscription_service import activate_paid_cycle
        activate_paid_cycle(first_company["id"], "test")

        created = self.client.post(
            "/companies",
            headers=first_headers,
            json={"company_name": "Workspace Two"},
        )
        self.assertEqual(created.status_code, 201)
        second_company = created.json()["company"]

        self.client.put(
            f"/companies/{first_company['id']}",
            headers=first_headers,
            json={"company_profile": "Private profile for workspace one."},
        )
        switched = self.client.post(
            "/auth/select-company",
            headers=first_headers,
            json={"company_id": second_company["id"]},
        )
        self.assertEqual(switched.status_code, 200)
        second_headers = {"Authorization": f"Bearer {switched.json()['token']}"}
        self.client.put(
            f"/companies/{second_company['id']}",
            headers=second_headers,
            json={"company_profile": "Private profile for workspace two."},
        )

        active = self.client.get("/auth/me", headers=second_headers)
        self.assertEqual(active.status_code, 200)
        self.assertEqual(active.json()["company"]["id"], second_company["id"])
        self.assertEqual(
            active.json()["company"]["company_profile"],
            "Private profile for workspace two.",
        )
        first = self.client.get(
            f"/companies/{first_company['id']}",
            headers=second_headers,
        )
        self.assertEqual(first.status_code, 200)
        self.assertEqual(
            first.json()["company_profile"],
            "Private profile for workspace one.",
        )

    def test_faq_chat_and_conversation_history(self):
        faq = self.client.post(
            "/faq",
            headers=self.headers,
            json={
                "question": "What are your business hours?",
                "answer": "We are open from 8 AM to 5 PM.",
            },
        )
        self.assertEqual(faq.status_code, 201)

        chat = self.client.post(
            "/chat",
            headers=self.headers,
            json={"text": "What are your business hours?", "language": "english"},
        )
        self.assertEqual(chat.status_code, 200)
        session_id = chat.json()["session_id"]
        self.assertEqual(chat.json()["reply"], "We are open from 8 AM to 5 PM.")

        sessions = self.client.get("/conversations", headers=self.headers)
        self.assertEqual(sessions.status_code, 200)
        self.assertTrue(any(item["session_id"] == session_id for item in sessions.json()))

        history = self.client.get(
            f"/conversations/{session_id}",
            headers=self.headers,
        )
        self.assertEqual(history.status_code, 200)
        self.assertEqual(
            [message["role"] for message in history.json()["messages"]],
            ["user", "assistant"],
        )

        deleted = self.client.delete(
            f"/conversations/{session_id}",
            headers=self.headers,
        )
        self.assertEqual(deleted.status_code, 200)
        self.assertTrue(deleted.json()["deleted"])
        missing = self.client.get(
            f"/conversations/{session_id}",
            headers=self.headers,
        )
        self.assertEqual(missing.status_code, 404)

    def test_protected_history_rejects_guests(self):
        response = self.client.get("/conversations")
        self.assertEqual(response.status_code, 401)

    def test_knowledge_document_upload_search_and_delete(self):
        uploaded = self.client.post(
            "/knowledge/documents",
            headers=self.headers,
            files={
                "file": (
                    "returns.md",
                    b"Returns Policy: Customers may return unused products within 45 days.",
                    "text/markdown",
                )
            },
        )
        self.assertEqual(uploaded.status_code, 201)
        document_id = uploaded.json()["id"]
        self.assertGreater(uploaded.json()["chunk_count"], 0)

        results = self.client.get(
            "/knowledge/search",
            headers=self.headers,
            params={"q": "unused products return"},
        )
        self.assertEqual(results.status_code, 200)
        self.assertEqual(results.json()[0]["document_id"], document_id)

        deleted = self.client.delete(
            f"/knowledge/documents/{document_id}", headers=self.headers
        )
        self.assertEqual(deleted.status_code, 200)
        self.assertTrue(deleted.json()["deleted"])

        from docx import Document
        from openpyxl import Workbook

        docx_buffer = BytesIO()
        document = Document()
        document.add_heading("Warranty Policy", 1)
        document.add_paragraph("All premium devices include a two-year warranty.")
        document.save(docx_buffer)
        docx_upload = self.client.post(
            "/knowledge/documents",
            headers=self.headers,
            files={
                "file": (
                    "warranty.docx",
                    docx_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(docx_upload.status_code, 201)
        self.assertEqual(docx_upload.json()["extraction_method"], "docx")

        xlsx_buffer = BytesIO()
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Products"
        worksheet.append(["Product", "Price"])
        worksheet.append(["Aura Premium", 499])
        workbook.save(xlsx_buffer)
        xlsx_upload = self.client.post(
            "/knowledge/documents",
            headers=self.headers,
            files={
                "file": (
                    "products.xlsx",
                    xlsx_buffer.getvalue(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(xlsx_upload.status_code, 201)
        self.assertEqual(xlsx_upload.json()["extraction_method"], "xlsx")

        blocked = self.client.post(
            "/knowledge/documents",
            headers=self.headers,
            files={
                "file": (
                    "dangerous.xlsm",
                    b"not-a-real-workbook",
                    "application/vnd.ms-excel.sheet.macroEnabled.12",
                )
            },
        )
        self.assertEqual(blocked.status_code, 415)

    def test_public_widget_is_company_scoped_and_token_can_rotate(self):
        settings = self.client.get("/widget/settings", headers=self.headers)
        self.assertEqual(settings.status_code, 200)
        token = settings.json()["public_token"]

        updated = self.client.put(
            "/widget/settings",
            headers=self.headers,
            json={
                "enabled": True,
                "title": "Core Flow Support",
                "welcome_message": "Welcome to our test widget.",
                "primary_color": "#123abc",
                "position": "left",
                "allowed_domains": ["allowed.example"],
            },
        )
        self.assertEqual(updated.status_code, 200)

        denied = self.client.get(
            f"/widget/public/{token}/config",
            headers={"Origin": "https://denied.example"},
        )
        self.assertEqual(denied.status_code, 403)

        config = self.client.get(
            f"/widget/public/{token}/config",
            headers={"Origin": "https://allowed.example"},
        )
        self.assertEqual(config.status_code, 200)
        self.assertEqual(config.json()["title"], "Core Flow Support")
        self.assertNotIn("company_id", config.json())

        chat = self.client.post(
            f"/widget/public/{token}/chat",
            json={"text": "What are your business hours?", "language": "english"},
            headers={"Origin": "https://allowed.example"},
        )
        self.assertEqual(chat.status_code, 200)
        self.assertEqual(chat.json()["reply"], "We are open from 8 AM to 5 PM.")

        analytics = self.client.get("/widget/analytics", headers=self.headers)
        self.assertEqual(analytics.status_code, 200)
        self.assertGreaterEqual(analytics.json()["widget_loads"], 1)
        self.assertGreaterEqual(analytics.json()["messages"], 1)

        session_id = chat.json()["session_id"]
        takeover = self.client.patch(
            f"/conversations/{session_id}/support",
            headers=self.headers,
            json={"mode": "human", "status": "open", "customer_name": "Widget Visitor"},
        )
        self.assertEqual(takeover.status_code, 200)
        self.assertEqual(takeover.json()["mode"], "human")

        queued = self.client.post(
            f"/widget/public/{token}/chat",
            json={"text": "I need a person.", "session_id": session_id},
            headers={"Origin": "https://allowed.example"},
        )
        self.assertEqual(queued.status_code, 200)
        self.assertEqual(queued.json()["status"], "waiting_for_agent")

        agent_reply = self.client.post(
            f"/conversations/{session_id}/reply",
            headers=self.headers,
            json={"text": "Hello, a human agent is here."},
        )
        self.assertEqual(agent_reply.status_code, 201)

        inbox = self.client.get("/conversations/support/inbox", headers=self.headers)
        self.assertEqual(inbox.status_code, 200)
        self.assertTrue(any(row["session_id"] == session_id for row in inbox.json()))

        rotated = self.client.post("/widget/token/rotate", headers=self.headers)
        self.assertEqual(rotated.status_code, 200)
        self.assertNotEqual(rotated.json()["public_token"], token)
        self.assertEqual(self.client.get(f"/widget/public/{token}/config").status_code, 404)

    def test_ai_provider_status_never_exposes_keys(self):
        response = self.client.get("/ai/providers", headers=self.headers)
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertTrue(data["managed_by_gateway"])
        self.assertNotIn("providers", data)
        self.assertNotIn("model", data)

        subscription = self.client.get("/subscription/status", headers=self.headers).json()
        for item in subscription.get("recent_ai_usage", []):
            self.assertNotIn("provider", item)
            self.assertNotIn("model", item)
            self.assertNotIn("estimated_cost", item)
        for item in subscription.get("recent_gateway_requests", []):
            self.assertNotIn("provider", item)
            self.assertNotIn("model", item)
            self.assertNotIn("provider_cost_usd", item)

    def test_facebook_webhook_verification_and_signature(self):
        os.environ["FACEBOOK_VERIFY_TOKEN"] = "aura-facebook-test-token"
        os.environ["FACEBOOK_APP_SECRET"] = "facebook-test-app-secret"
        verified = self.client.get(
            "/webhooks/facebook",
            params={
                "hub.mode": "subscribe",
                "hub.verify_token": "aura-facebook-test-token",
                "hub.challenge": "challenge-123",
            },
        )
        self.assertEqual(verified.status_code, 200)
        self.assertEqual(verified.text, "challenge-123")

        denied = self.client.get(
            "/webhooks/facebook",
            params={
                "hub.mode": "subscribe",
                "hub.verify_token": "wrong-token",
                "hub.challenge": "challenge-123",
            },
        )
        self.assertEqual(denied.status_code, 403)

        payload = {
            "object": "page",
            "entry": [{
                "id": "page-1",
                "messaging": [{
                    "sender": {"id": "customer-1"},
                    "message": {"mid": "message-1", "text": "Hello Aura"},
                }],
            }],
        }
        body = json.dumps(payload, separators=(",", ":")).encode("utf-8")
        digest = hmac.new(
            b"facebook-test-app-secret", body, hashlib.sha256
        ).hexdigest()
        with patch(
            "app.services.facebook_messenger_service.process_webhook"
        ) as processor:
            processor.return_value = None
            accepted = self.client.post(
                "/webhooks/facebook",
                content=body,
                headers={
                    "content-type": "application/json",
                    "x-hub-signature-256": f"sha256={digest}",
                },
            )
        self.assertEqual(accepted.status_code, 200)
        self.assertEqual(accepted.json()["status"], "received")
        processor.assert_called_once()

        rejected = self.client.post(
            "/webhooks/facebook",
            content=body,
            headers={
                "content-type": "application/json",
                "x-hub-signature-256": "sha256=invalid",
            },
        )
        self.assertEqual(rejected.status_code, 403)

    def test_facebook_message_routes_through_aura_and_replies(self):
        from app.services import facebook_messenger_service

        company_id = self.client.get(
            "/auth/me", headers=self.headers
        ).json()["company"]["id"]
        os.environ["FACEBOOK_COMPANY_ID"] = str(company_id)
        os.environ["FACEBOOK_AI_PROVIDERS"] = "gemini"
        payload = {
            "object": "page",
            "entry": [{
                "id": "page-route-test",
                "messaging": [{
                    "sender": {"id": "customer-route-test"},
                    "message": {
                        "mid": "unique-route-message",
                        "text": "What are your business hours?",
                    },
                }],
            }],
        }
        with patch(
            "app.services.facebook_messenger_service.handle_chat",
            return_value={"reply": "We are open from 8 AM to 5 PM."},
        ) as chat_handler, patch(
            "app.services.facebook_messenger_service.send_text"
        ) as sender:
            asyncio.run(facebook_messenger_service.process_webhook(payload))
        chat_handler.assert_called_once()
        self.assertEqual(
            chat_handler.call_args.kwargs["session_id"],
            "facebook:page-route-test:customer-route-test",
        )
        sender.assert_called_once_with(
            "customer-route-test", "We are open from 8 AM to 5 PM."
        )

    def test_simple_facebook_connect_is_encrypted_and_company_scoped(self):
        from app.database.connection import get_connection

        disconnected = self.client.get(
            "/webhooks/facebook/connection", headers=self.headers
        )
        self.assertEqual(disconnected.status_code, 200)
        self.assertFalse(disconnected.json()["connected"])

        with patch(
            "app.services.facebook_messenger_service.connect_page",
            return_value={"page_id": "simple-page-123", "page_name": "Simple Test Page"},
        ):
            connected = self.client.post(
                "/webhooks/facebook/connection",
                headers=self.headers,
                json={
                    "page_access_token": "test-page-token-that-is-long-enough",
                    "app_secret": "test-app-secret",
                },
            )
        self.assertEqual(connected.status_code, 200, connected.text)
        data = connected.json()
        self.assertTrue(data["connected"])
        self.assertEqual(data["page_name"], "Simple Test Page")
        self.assertNotIn("page_access_token", data)
        self.assertNotIn("app_secret", data)

        row = get_connection().cursor().execute(
            "SELECT encrypted_page_token, encrypted_app_secret FROM facebook_connections WHERE page_id = ?",
            ("simple-page-123",),
        ).fetchone()
        self.assertIsNotNone(row)
        self.assertNotIn("test-page-token", row["encrypted_page_token"])
        self.assertNotIn("test-app-secret", row["encrypted_app_secret"])

        verified = self.client.get(
            "/webhooks/facebook",
            params={
                "hub.mode": "subscribe",
                "hub.verify_token": data["verify_token"],
                "hub.challenge": "simple-connect-ok",
            },
        )
        self.assertEqual(verified.status_code, 200)
        self.assertEqual(verified.text, "simple-connect-ok")

        removed = self.client.delete(
            "/webhooks/facebook/connection", headers=self.headers
        )
        self.assertEqual(removed.status_code, 200)
        self.assertTrue(removed.json()["disconnected"])

    def test_one_click_facebook_oauth_page_selection(self):
        os.environ["FACEBOOK_APP_ID"] = "test-meta-app-id"
        os.environ["FACEBOOK_APP_SECRET"] = "test-meta-app-secret"
        started = self.client.get(
            "/webhooks/facebook/oauth/start", headers=self.headers
        )
        self.assertEqual(started.status_code, 200, started.text)
        authorization_url = started.json()["authorization_url"]
        query = parse_qs(urlparse(authorization_url).query)
        self.assertEqual(query["client_id"], ["test-meta-app-id"])
        self.assertIn("pages_show_list", query["scope"][0])
        self.assertNotIn("pages_read_engagement", query["scope"][0])

        with patch(
            "app.services.facebook_messenger_service.exchange_oauth_code",
            return_value="oauth-user-token",
        ), patch(
            "app.services.facebook_messenger_service.list_user_pages",
            return_value=[{
                "page_id": "oauth-page-1",
                "page_name": "OAuth Test Page",
                "page_token": "oauth-page-token-that-is-long-enough",
            }],
        ):
            callback = self.client.get(
                "/webhooks/facebook/oauth/callback",
                params={"code": "test-code", "state": query["state"][0]},
            )
        self.assertEqual(callback.status_code, 200, callback.text)
        match = re.search(r"session:([^}]+)", callback.text)
        self.assertIsNotNone(match)
        session = json.loads(match.group(1))

        pages = self.client.get(
            "/webhooks/facebook/oauth/pages",
            headers=self.headers,
            params={"session": session},
        )
        self.assertEqual(pages.status_code, 200, pages.text)
        self.assertEqual(pages.json()["pages"][0]["page_name"], "OAuth Test Page")
        self.assertNotIn("page_token", pages.json()["pages"][0])

        with patch(
            "app.services.facebook_messenger_service.connect_page",
            side_effect=RuntimeError("Facebook rejected the Page permission."),
        ):
            failed = self.client.post(
                "/webhooks/facebook/oauth/complete/oauth-page-1",
                headers=self.headers,
                params={"session": session},
            )
        self.assertEqual(failed.status_code, 400, failed.text)
        self.assertIn("Facebook rejected", failed.json()["detail"])

        with patch(
            "app.services.facebook_messenger_service.connect_page",
            return_value={"page_id": "oauth-page-1", "page_name": "OAuth Test Page"},
        ):
            completed = self.client.post(
                "/webhooks/facebook/oauth/complete/oauth-page-1",
                headers=self.headers,
                params={"session": session},
            )
        self.assertEqual(completed.status_code, 200, completed.text)
        self.assertTrue(completed.json()["connected"])

        self.client.delete("/webhooks/facebook/connection", headers=self.headers)

    def test_customer_cannot_manage_server_provider_keys(self):
        connected = self.client.post(
            "/ai/providers/openai/connect",
            headers=self.headers,
            json={"api_key": "sk-customer-secret-example-1234"},
        )
        self.assertEqual(connected.status_code, 403)
        self.assertNotIn("sk-customer", connected.text)
        self.assertEqual(
            self.client.delete("/ai/providers/openai", headers=self.headers).status_code,
            403,
        )

    def test_gateway_ignores_customer_provider_selection(self):
        from app.services import ai_gateway, ai_service

        os.environ["OPENAI_API_KEY"] = "test-openai-key"
        ai_gateway.AI_MODEL_COSTS_USD["test-core-model"] = {
            "input_per_million": 1.0,
            "output_per_million": 1.0,
        }
        try:
            with patch("app.services.ai_service.get_server_model", return_value="test-core-model"), patch(
                "app.services.ai_service.generate_reply",
                side_effect=lambda prompt, provider, *args: ai_service.AIResult(
                    f"Answer from {provider}", 20, 10
                ),
            ) as generated:
                response = self.client.post(
                    "/chat",
                    headers=self.headers,
                    json={
                        "text": "Compare two launch ideas.",
                        "language": "english",
                        "providers": ["openai", "deepseek"],
                    },
                )
            self.assertEqual(response.status_code, 200)
            data = response.json()
            self.assertEqual(len(data["provider_responses"]), 1)
            self.assertEqual(generated.call_args.args[1], "openai")
            self.assertIn("Answer from openai", data["reply"])
            self.assertNotIn("provider", data["provider_responses"][0])
            self.assertNotIn("model", data["provider_responses"][0])
        finally:
            os.environ["OPENAI_API_KEY"] = ""
            ai_gateway.AI_MODEL_COSTS_USD.pop("test-core-model", None)

    def test_company_profile_html_is_cleaned_and_grounds_ai_prompt(self):
        from app.services.chat_service import _build_prompt

        prompt = _build_prompt(
            "What is the return policy?",
            '<p><img src="https://noise.example/image.jpg">Returns are accepted within 30 days.</p>',
            "english",
        )
        self.assertIn("Returns are accepted within 30 days.", prompt)
        self.assertNotIn("<img", prompt)
        self.assertNotIn("noise.example", prompt)
        self.assertIn("Do not invent company facts", prompt)

    def test_z_paid_subscription_expiration_and_renewal(self):
        company_id = self.client.get("/auth/me", headers=self.headers).json()["company"]["id"]
        from app.services.subscription_service import activate_paid_cycle
        activated = activate_paid_cycle(company_id, "test")
        self.assertIsNotNone(activated["expires_at"])

        from app.database.connection import get_connection

        conn = get_connection()
        conn.cursor().execute(
            """
            UPDATE subscriptions
            SET expires_at = '2020-01-01T00:00:00+00:00'
            WHERE company_id = ?
            """,
            (company_id,),
        )
        conn.commit()

        status = self.client.get("/subscription/status", headers=self.headers)
        self.assertEqual(status.status_code, 200)
        self.assertEqual(status.json()["subscription"]["status"], "expired")
        self.assertEqual(status.json()["notification"]["level"], "danger")

        blocked = self.client.post(
            "/chat",
            headers=self.headers,
            json={"text": "This should be blocked while expired."},
        )
        self.assertIn("expired", blocked.json()["reply"].lower())

        renewed = activate_paid_cycle(company_id, "test")
        self.assertEqual(renewed["status"], "active")

    def test_zz_verified_account_email_change(self):
        requested = self.client.post(
            "/auth/account/email/request",
            headers=self.headers,
            json={"new_email": "updated-core-flow@example.com"},
        )
        self.assertEqual(requested.status_code, 200)
        code = requested.json()["development_code"]
        verified = self.client.post(
            "/auth/account/email/verify",
            headers=self.headers,
            json={
                "new_email": "updated-core-flow@example.com",
                "code": code,
            },
        )
        self.assertEqual(verified.status_code, 200)
        self.assertEqual(
            verified.json()["user"]["email"],
            "updated-core-flow@example.com",
        )

    def test_passwordless_email_login_for_new_user(self):
        email = "passwordless@yahoo.com"
        requested = self.client.post("/auth/request-code", json={"email": email})
        self.assertEqual(requested.status_code, 200)
        self.assertEqual(requested.json()["delivery"], "development")
        code = requested.json()["development_code"]

        wrong = self.client.post(
            "/auth/verify-code",
            json={"email": email, "code": "000000", "company_name": "Yahoo Company"},
        )
        self.assertEqual(wrong.status_code, 400)

        verified = self.client.post(
            "/auth/verify-code",
            json={"email": email, "code": code, "company_name": "Yahoo Company"},
        )
        self.assertEqual(verified.status_code, 200)
        self.assertTrue(verified.json()["new_user"])
        self.assertEqual(verified.json()["company"]["company_name"], "Yahoo Company")

        reused = self.client.post(
            "/auth/verify-code",
            json={"email": email, "code": code},
        )
        self.assertEqual(reused.status_code, 400)

    def test_passwordless_email_login_for_existing_user(self):
        requested = self.client.post(
            "/auth/request-code",
            json={"email": "core-flow@example.com"},
        )
        code = requested.json()["development_code"]
        verified = self.client.post(
            "/auth/verify-code",
            json={"email": "core-flow@example.com", "code": code},
        )
        self.assertEqual(verified.status_code, 200)
        self.assertFalse(verified.json()["new_user"])


if __name__ == "__main__":
    unittest.main()

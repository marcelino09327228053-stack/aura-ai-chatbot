"""Free, local text extraction for common business documents and images."""

import csv
import html
import re
from io import BytesIO, StringIO
from pathlib import Path

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".csv", ".html", ".htm",
    ".pdf", ".docx", ".xlsx", ".pptx",
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
MAX_PDF_PAGES = 30
_ocr_engine = None


class ExtractionError(ValueError):
    pass


def extract_text(name: str, raw: bytes) -> tuple[str, str]:
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            "Supported files: PDF, DOCX, XLSX, PPTX, TXT, Markdown, CSV, HTML, "
            "PNG, JPG, WebP, BMP, and TIFF."
        )
    try:
        if extension in {".txt", ".md", ".csv", ".html", ".htm"}:
            return _extract_plain(extension, raw), "text"
        if extension == ".pdf":
            return _extract_pdf(raw)
        if extension == ".docx":
            return _extract_docx(raw), "docx"
        if extension == ".xlsx":
            return _extract_xlsx(raw), "xlsx"
        if extension == ".pptx":
            return _extract_pptx(raw), "pptx"
        if extension in IMAGE_EXTENSIONS:
            return _extract_image(raw), "ocr"
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not read {extension or 'this file'}: {exc}") from exc
    raise ExtractionError("Unsupported document.")


def _extract_plain(extension: str, raw: bytes) -> str:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    if extension == ".csv":
        rows = csv.reader(StringIO(text))
        text = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    elif extension in {".html", ".htm"}:
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        text = html.unescape(text)
    return _clean(text)


def _extract_pdf(raw: bytes) -> tuple[str, str]:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(raw))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ExtractionError(f"PDF is limited to {MAX_PDF_PAGES} pages per upload.")
    native = "\n".join((page.extract_text() or "") for page in reader.pages)
    native = _clean(native)
    if len(native) >= 80:
        return native, "pdf-text"

    import fitz

    document = fitz.open(stream=raw, filetype="pdf")
    recognized = []
    for page in document:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1.6, 1.6), alpha=False)
        recognized.append(_extract_image(pixmap.tobytes("png")))
    return _clean("\n".join(recognized)), "pdf-ocr"


def _extract_docx(raw: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(raw))
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text for cell in row.cells))
    return _clean("\n".join(blocks))


def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"Sheet: {sheet.title}")
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number > 5000:
                lines.append("[Sheet truncated after 5000 rows]")
                break
            values = [str(value) if value is not None else "" for value in row]
            if any(values):
                lines.append(" | ".join(values))
    workbook.close()
    return _clean("\n".join(lines))


def _extract_pptx(raw: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(raw))
    lines = []
    for number, slide in enumerate(presentation.slides, start=1):
        lines.append(f"Slide {number}")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                lines.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
    return _clean("\n".join(lines))


def _extract_image(raw: bytes) -> str:
    global _ocr_engine
    import numpy as np
    from PIL import Image
    from rapidocr_onnxruntime import RapidOCR

    image = Image.open(BytesIO(raw)).convert("RGB")
    if image.width * image.height > 25_000_000:
        image.thumbnail((5000, 5000))
    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    result, _ = _ocr_engine(np.asarray(image))
    if not result:
        return ""
    return _clean("\n".join(str(item[1]) for item in result if len(item) >= 2))


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
